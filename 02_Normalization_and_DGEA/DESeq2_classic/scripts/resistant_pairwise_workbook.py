"""2026-10-19: workbook of the direct resistant-genotype contrasts (DGEA_resistant_pairwise.R): one sheet per genotype pair,
0/6/24 hpi side by side, numeric cells, method note in the header rows. Run from plant_immunity_Vitis/:
   python3 02_Normalization_and_DGEA/DESeq2_classic/scripts/resistant_pairwise_workbook.py <out.xlsx>"""
import sys, csv, openpyxl
from openpyxl.styles import Font, Alignment
B = "02_Normalization_and_DGEA"; R = f"{B}/DESeq2_classic/results/DGEA_reanalysis"
ann = {r["PN40024_genotype_ENSMBL_ID"]: (r["Athaliana_TAIR10_homolog"], r["Athal.common.NAme"])
       for r in csv.DictReader(open("data_files/26169genes_with_AthalHomologs_allIDs_exprPatterns_TAIR10ids.csv"), delimiter="\t")}
audit = {(r["contrast"], int(r["time_hpi"])): r for r in csv.DictReader(open(f"{R}/resistant_pairwise_batch_audit.csv"))}
summ = {(r["contrast"], int(r["time_hpi"])): r for r in csv.DictReader(open(f"{R}/resistant_pairwise_DE_summary.csv"))}
num = lambda v: None if v in ("", "NA") else (v == "TRUE" if v in ("TRUE", "FALSE") else float(v))
wb = openpyxl.Workbook(); wb.remove(wb.active)
for pr in ("Rpv12_vs_Rpv12+1", "Rpv12_vs_Rpv12+1+3", "Rpv12+1_vs_Rpv12+1+3"):
    g1, g2 = pr.split("_vs_"); data = {}
    for t in (0, 6, 24):
        for r in csv.DictReader(open(f"{B}/{pr}/DGEA_{pr}_{t}hpi_deseq2_NB_BE_model.csv")):
            d = data.setdefault(r["gene"], {})
            for k in ("log2FC_apeglm", "lfcSE_apeglm", "pvalue", "FDR", "DE_primary", "DE_both_directions", "zero_counts_first", "zero_counts_second", "numerically_unstable"):
                d[(k, t)] = num(r[k])
    ws = wb.create_sheet(f"{g1} vs {g2}")
    batches = "; ".join(f"{t} hpi: {g1} {audit[(pr, t)]['first_batches']}, {g2} {audit[(pr, t)]['second_batches']}" for t in (0, 6, 24))
    counts = "; ".join(f"{t} hpi: {summ[(pr, t)]['DE']} DE ({summ[(pr, t)]['up_first']} higher, {summ[(pr, t)]['down_first']} lower in {g1}), "
                       f"{summ[(pr, t)]['DE_both_directions']} DE in both directions" for t in (0, 6, 24))
    notes = [f"Direct comparison {g1} vs {g2} at 0, 6 and 24 hpi; log2FC = {g1} / {g2} (positive = higher in {g1}).",
             "DESeq2 on raw counts of all 36 libraries (row sum >= 15; 26,169 genes), design ~ batch + condition (condition = genotype x time; "
             "additive sequencing-batch covariate), Wald test, apeglm-shrunken log2FC, BH FDR within each comparison. "
             "DE = FDR < 0.05 and |log2FC| > 1 (same model and rule as the comparisons with the susceptible control).",
             "DE_both = also called when the reference genotype is swapped (apeglm shrinkage is estimated per tested coefficient, so a few "
             "borderline calls depend on the direction of the comparison). zero_first / zero_second = no counts in that genotype at that time; "
             "p and FDR are NA where both are zero (log2FC not estimable).",
             f"Sequencing batches - {batches}.", f"Counts - {counts}."]
    for i, n in enumerate(notes, 1):
        ws.cell(i, 1, n).alignment = Alignment(wrap_text=False)
    hdr = ["gene", "Athal_homolog", "Athal_name"] + [f"{k}_{t}hpi" for t in (0, 6, 24) for k in
           ("log2FC", "lfcSE", "p", "FDR", "DE", "DE_both", "zero_first", "zero_second", "unstable")]
    hr = len(notes) + 2
    for j, h in enumerate(hdr, 1): ws.cell(hr, j, h).font = Font(bold=True)
    keys = ("log2FC_apeglm", "lfcSE_apeglm", "pvalue", "FDR", "DE_primary", "DE_both_directions", "zero_counts_first", "zero_counts_second", "numerically_unstable")
    for i, gene in enumerate(sorted(data), hr + 1):
        a = ann.get(gene, ("", "")); row = [gene, a[0] or None, a[1] or None] + [data[gene].get((k, t)) for t in (0, 6, 24) for k in keys]
        for j, v in enumerate(row, 1): ws.cell(i, j, v)
    ws.freeze_panes = ws.cell(hr + 1, 2)
wb.save(sys.argv[1]); print("saved", sys.argv[1])
