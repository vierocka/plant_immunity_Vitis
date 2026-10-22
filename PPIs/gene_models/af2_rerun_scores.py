#!/usr/bin/env python3
# Table S10 sheets 13-16 (gene-model comparison and repeated AF2 predictions). Inputs/outputs outside the repository
# (reference annotations, AF2 run, audit results) are read from VITIS_BASE.
# Scores the AlphaFold2-Multimer re-runs with the best available gene models (2026-10-22), using the screen's definitions
# (reproduced exactly on the original SOBIR1-LRIP1 run): per model ipTM, pTM, mean pLDDT (scores JSON), interface contacts
# = inter-chain CA-CA < 8 A with pLDDT >= 50 in both residues (PDB), contact-filtered mean PAE = mean of PAE[chain A residue,
# chain B residue] over contacts, fraction of contacts with PAE < 10; single-model pass = ipTM >= 0.70, >= 5 contacts and
# (contact PAE < 10 or >= 50% of contacts with PAE < 10); strict = all 5 models ipTM >= 0.75 and >= 5 contacts, and the mean
# contact-filtered PAE across the five models < 10. Original values (PN40024 v4 models) from Table S10, all_positive_hits.
# Usage (from the folder containing plant_immunity_Vitis): python3 plant_immunity_Vitis/PPIs/gene_models/af2_rerun_scores.py
import glob, gzip, json, math, re, os, csv, openpyxl
B = os.environ.get("VITIS_BASE", os.path.abspath(os.path.join(os.path.dirname(__file__), "../../..")) + "/")   # folder that contains plant_immunity_Vitis/
D = B + "gene_model_audit/corrected_models/AF2/"
OUT = B + "gene_model_audit/results/af2_rerun_scores.tsv"
op = lambda p: gzip.open(p, "rt") if p.endswith(".gz") else open(p)

def score_model(pdb, js):
    ca = {}
    for l in op(pdb):
        if l.startswith("ATOM") and l[12:16].strip() == "CA":
            ca.setdefault(l[21], {})[int(l[22:26])] = (float(l[30:38]), float(l[38:46]), float(l[46:54]), float(l[60:66]))
    A, Bc = sorted(ca)[:2]; LA = len(ca[A])
    s = json.load(op(js)); pae = s["pae"]
    c = [(i, j) for i, a in ca[A].items() for j, b in ca[Bc].items()
         if a[3] >= 50 and b[3] >= 50 and math.dist(a[:3], b[:3]) < 8]
    p = [pae[i - 1][LA + j - 1] for i, j in c]
    cp = sum(p) / len(p) if p else float("nan"); fr = sum(x < 10 for x in p) / len(p) if p else 0.0
    single = s["iptm"] >= 0.70 and len(c) >= 5 and (cp < 10 or fr >= 0.5)
    return dict(iptm=s["iptm"], ptm=s["ptm"], plddt=sum(s["plddt"]) / len(s["plddt"]), contacts=len(c), cpae=cp, frac10=fr,
                single=single, lenA=LA, lenB=len(ca[Bc]), contact_pairs=c, pae=pae)

def score_job(prefix):
    ms = []
    for k in range(1, 6):
        pdb = glob.glob(f"{prefix}_unrelaxed_rank_{k:03d}_*.pdb*"); js = glob.glob(f"{prefix}_scores_rank_{k:03d}_*.json*")
        if pdb and js: m = score_model(pdb[0], js[0]); m["rank"] = k; m["model"] = re.search(r"model_(\d)", pdb[0]).group(1); ms.append(m)
    if len(ms) < 5: return ms, "incomplete"
    strict = all(m["iptm"] >= 0.75 and m["contacts"] >= 5 for m in ms) and sum(m["cpae"] for m in ms) / 5 < 10
    return ms, "PASS" if strict else "FAIL"

# original v4 results from Table S10
ws = list(openpyxl.load_workbook(B + "plant_immunity_Vitis/PPIs/results/Supplementary_Table_8.xlsx", read_only=True)["all_positive_hits"].iter_rows(values_only=True))
h = ws[4]; ix = {k: h.index(k) for k in ("geneA_gene_id", "geneB_gene_id", "ipTM", "contact_filtered_mean_iPAE", "interface_contacts", "Strict_5model_verdict", "rank")}
orig, sec = {}, None
for r in ws[5:]:
    if r[3] and str(r[3]).startswith("SECTION"): sec = str(r[3]); continue
    if sec and "CANDIDATE" in sec and r[ix["geneA_gene_id"]]:
        orig.setdefault(frozenset((r[ix["geneA_gene_id"]], r[ix["geneB_gene_id"]])), []).append(r)

jobs = sorted({re.sub(r"(_scores_rank.*|_unrelaxed_rank.*|\.done\.txt.*)$", "", os.path.basename(p)) for p in glob.glob(D + "*_rank_*")})
rows = []
for j in jobs:
    ms, verdict = score_job(D + j)
    g = re.findall(r"Vitvi\d\dg\d{5}", j); o = orig.get(frozenset(g), [])
    oi = [float(r[ix["ipTM"]]) for r in o]; op_ = [float(r[ix["contact_filtered_mean_iPAE"]]) for r in o]; oc = [int(r[ix["interface_contacts"]]) for r in o]
    rows.append(dict(job=j, geneA=g[0], geneB=g[1], lenA=ms[0]["lenA"] if ms else "", lenB=ms[0]["lenB"] if ms else "", n_models=len(ms),
                     iptm_min=min(m["iptm"] for m in ms), iptm_max=max(m["iptm"] for m in ms), iptm_mean=sum(m["iptm"] for m in ms) / len(ms),
                     ptm_mean=sum(m["ptm"] for m in ms) / len(ms), plddt_mean=sum(m["plddt"] for m in ms) / len(ms),
                     contacts_min=min(m["contacts"] for m in ms), contacts_max=max(m["contacts"] for m in ms),
                     cpae_mean=sum(m["cpae"] for m in ms) / len(ms), cpae_min=min(m["cpae"] for m in ms), cpae_max=max(m["cpae"] for m in ms),
                     models_single_pass=sum(m["single"] for m in ms), strict=verdict,
                     v4_iptm_range=f"{min(oi):.2f}-{max(oi):.2f}" if oi else "", v4_cpae_mean=round(sum(op_) / len(op_), 2) if op_ else "",
                     v4_contacts_range=f"{min(oc)}-{max(oc)}" if oc else "", v4_strict=o[0][ix["Strict_5model_verdict"]] if o else ""))
with open(OUT, "w") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter="\t"); w.writeheader(); w.writerows(rows)
for r in rows:
    print(f"{r['job']:42s} {r['lenA']:>4}+{r['lenB']:<4} ipTM {r['iptm_min']:.2f}-{r['iptm_max']:.2f} contacts {r['contacts_min']}-{r['contacts_max']} "
          f"cPAE {r['cpae_mean']:.2f} ({r['cpae_min']:.1f}-{r['cpae_max']:.1f}) {r['strict']:4s} | v4: ipTM {r['v4_iptm_range']} cPAE {r['v4_cpae_mean']} {r['v4_strict']}")

# ---- Table S10 sheet "AF2_rerun_gene_models" (repo copy + BMC copy) ----
import shutil
from openpyxl.styles import Font
REPO = B + "plant_immunity_Vitis/PPIs/results/Supplementary_Table_8.xlsx"
BMC = B + "draft/Sections_by_VK/BMC_PlantBiology/Supplementary_tables/Additional_file_18_TableS10.xlsx"
used = {r["gene"]: r for r in csv.DictReader(open(B + "gene_model_audit/corrected_models/network_proteins_best_models.tsv"), delimiter="\t")}
wb = openpyxl.load_workbook(REPO)
if "AF2_rerun_gene_models" in wb.sheetnames: del wb["AF2_rerun_gene_models"]
w = wb.create_sheet("AF2_rerun_gene_models")
w.append(["The ten candidate pairs of Figure 4A predicted again with the longest available gene model of each protein (sheet gene_models_network; "
          "PN40024 v4 model kept for SOBIR1 and MRH1), with the same ColabFold settings (v1.6.3; alphafold2_multimer_v3; 5 models; 3 recycles; "
          "MMseqs2 MSA) and the same scoring as the screen: interface contacts = inter-chain CA-CA < 8 A with pLDDT >= 50 in both residues; "
          "contact-filtered PAE = mean PAE (chain A residue -> chain B residue) over contacts; single-model criterion (passes_single_model) = "
          "ipTM >= 0.70, >= 5 contacts, and either contact-filtered PAE < 10 A or >= 50% of contacts with PAE < 10 A (so a model with a mean "
          "PAE just above 10 A can pass); strict = ipTM >= 0.75 and >= 5 contacts in all five models and mean contact-filtered PAE < 10 A. "
          "Original = PN40024 v4 models (sheet all_positive_hits). The longer models are themselves unvalidated predictions (sheet "
          "gene_models_loci). Reference rows at the end: EDS1 complex pairs of the original screen in which both v4 models are complete."])
w["A1"].font = Font(italic=True); w.append([])
hdr = ["pair (chain A : chain B)", "model_A", "length_A_aa", "model_B", "length_B_aa", "ipTM_min", "ipTM_max", "ipTM_mean", "pTM_mean",
       "pLDDT_mean", "contacts_min", "contacts_max", "contact_PAE_mean_A", "models_passing_single_model_criterion", "strict_verdict",
       "original_ipTM_range", "original_contact_PAE_mean_A", "original_strict_verdict"]
w.append(hdr)
for c in w[3]: c.font = Font(bold=True)
src = lambda g: f'{used[g]["source"]} {re.sub(r"^lcl[|]NC_[0-9.]+_prot_", "", used[g]["model"]).rsplit("_", 1)[0] if used[g]["source"] == "RefSeq" else used[g]["model"]}'
for r in rows:
    w.append([r["job"], src(r["geneA"]), r["lenA"], src(r["geneB"]), r["lenB"], round(r["iptm_min"], 2), round(r["iptm_max"], 2), round(r["iptm_mean"], 3),
              round(r["ptm_mean"], 3), round(r["plddt_mean"], 1), r["contacts_min"], r["contacts_max"], round(r["cpae_mean"], 2),
              f'{r["models_single_pass"]}/5', r["strict"], r["v4_iptm_range"], r["v4_cpae_mean"], r["v4_strict"]])
w.append([]); w.append(["Per model"]); w.cell(w.max_row, 1).font = Font(bold=True)
w.append(["pair", "rank", "AF2 parameter set", "ipTM", "pTM", "mean_pLDDT", "interface_contacts", "contact_PAE_mean_A", "fraction_contacts_PAE_lt_10", "passes_single_model"])
for c in w[w.max_row]: c.font = Font(bold=True)
for j in jobs:
    ms, _ = score_job(D + j)
    for m in ms:
        w.append([j, m["rank"], int(m["model"]), round(m["iptm"], 3), round(m["ptm"], 3), round(m["plddt"], 1), m["contacts"],
                  round(m["cpae"], 2), round(m["frac10"], 3), m["single"]])
# reference: EDS1 x SAG101/PAD4 pairs of the original v4 screen in which both models are complete (sheet gene_models_43)
g43 = list(wb["gene_models_43"].iter_rows(values_only=True)); h43 = g43[2]
comp = {r[0]: r[1] for r in g43[3:] if r[0] and str(r[0]).startswith("Vitvi") and r[h43.index("v4_complete")] == "yes"}
pc = list(wb["positive_control_EDS1_SAG101_PA"].iter_rows(values_only=True)); hp = pc[4]
ip = {k: hp.index(k) for k in ("geneA_gene_id", "geneB_gene_id", "ipTM", "contact_filtered_mean_iPAE", "Strict_5model_verdict")}
ref = {}
for r in pc[5:]:
    a, b = r[ip["geneA_gene_id"]], r[ip["geneB_gene_id"]]
    if a in comp and b in comp and {"EDS1"} & {comp[a], comp[b]} and {"SAG101", "PAD4"} & {comp[a], comp[b]} and r[ip["ipTM"]] not in (None, ""):
        ref.setdefault((a, b), []).append(r)
w.append([]); w.append(["Reference: EDS1 complex pairs with complete v4 models (original screen, sheet positive_control_EDS1_SAG101_PA)"])
w.cell(w.max_row, 1).font = Font(bold=True)
w.append(["gene A", "name A", "gene B", "name B", "ipTM_min", "ipTM_max", "contact_PAE_mean_A", "strict_verdict"])
for c in w[w.max_row]: c.font = Font(bold=True)
for (a, b), rs in sorted(ref.items()):
    i_ = [float(r[ip["ipTM"]]) for r in rs]; p_ = [float(r[ip["contact_filtered_mean_iPAE"]]) for r in rs]
    w.append([a, comp[a], b, comp[b], round(min(i_), 2), round(max(i_), 2), round(sum(p_) / len(p_), 2), rs[0][ip["Strict_5model_verdict"]]])
print(len(ref), "reference pairs; ipTM", min(float(r[ip["ipTM"]]) for v in ref.values() for r in v), "-", max(float(r[ip["ipTM"]]) for v in ref.values() for r in v))
ORDER = ["all_positive_hits", "Domain_architecture", "domain_summary", "roadmap_all_PPI_tested_pairs", "within_module_all_models",
         "cross_module_all_models", "negative_controls_all_models", "positive_control_EDS1_SAG101_PA", "promiscuity_stickiness_tests_al",
         "failures_zero_contacts", "failures_chainA_empty", "failures_chainB_empty", "gene_models_network", "gene_models_43",
         "AF2_rerun_gene_models", "gene_models_loci"]   # sheet order of the Table S10 caption
wb._sheets.sort(key=lambda x: ORDER.index(x.title) if x.title in ORDER else 99)
wb.save(REPO); shutil.copyfile(REPO, BMC); print("sheet AF2_rerun_gene_models written (repo + BMC)")
