#!/usr/bin/env python3
# Table S10 sheets 13-16 (gene-model comparison and repeated AF2 predictions). Inputs/outputs outside the repository
# (reference annotations, AF2 run, audit results) are read from VITIS_BASE.
# Cross-annotation check of the 43 PN40024 v4 gene models (2026-10-19), stage 2 of PPIs/scripts/gene_model_audit.py.
# Each v4 locus is placed on the T2T (GCA_030704535.1) and 12X assemblies by genome alignment (minimap2 -x asm5),
# and every protein model overlapping it on the same strand is taken from three independent annotations:
#   T2T GenBank (VitviT2T_, WJZ proteins), NCBI RefSeq Gnomon on T2T (XP_/NP_), Phytozome v2.1 on 12X.
# Each model is compared with the v4 model's best Arabidopsis hit (same N/C "missing" rule as stage 1).
# Usage (from the folder containing plant_immunity_Vitis): python3 plant_immunity_Vitis/PPIs/gene_models/gene_model_cross_annotation.py \
#   plant_immunity_Vitis/PPIs/results/gene_model_audit_43genes.tsv > gene_model_audit/results/gene_model_cross_annotation_43genes.tsv
import sys, os, re, gzip, subprocess, tempfile, collections
B = os.environ.get("VITIS_BASE", os.path.abspath(os.path.join(os.path.dirname(__file__), "../../..")) + "/")   # folder that contains plant_immunity_Vitis/
R = B + "reference/Vitis_vinifera/"
V4_GENOME, V4_GFF, V4_PEP = R + "GG/reference.fa", R + "GG/Vitis_vinifera.PN40024.v4.56.gff3", R + "GG/Vitis_vinifera.PN40024.v4.pep.all.fa"
T2T = R + "ASM3070453v1/GCA_030704535.1/"
T2T_GENOME, T2T_GFF, T2T_PEP, RS_PEP = T2T + "GCA_030704535.1_ASM3070453v1_genomic.fna", T2T + "genomic.gff", T2T + "protein.faa", T2T + "GCF_030704535.1_ASM3070453v1_translated_cds.faa"
PZ = R + "phytozome/"
PZ_GENOME, PZ_GFF, PZ_PEP = PZ + "assembly/Vvinifera_457_Genoscope.12X.fa.gz", PZ + "annotation/Vvinifera_457_v2.1.gene.gff3", PZ + "annotation/Vvinifera_457_v2.1.protein.fa"
TAIR = R + "TAIR10/GCF_000001735.3_TAIR10_protein.faa"
TMP = tempfile.mkdtemp()
MISS = 25
MAXLEN = 1.25   # alternative model longer than 1.25 x the Arabidopsis homolog = likely fused (e.g. 2,145-aa SAG101-cluster model); not accepted

def fasta(path, keep=None, key=lambda h: h.split()[0]):
    seqs, name = {}, None
    op = gzip.open if path.endswith(".gz") else open
    for l in op(path, "rt"):
        if l.startswith(">"):
            name = key(l[1:]); seqs[name] = [] if keep is None or keep(name) else None
        elif name is not None and seqs.get(name) is not None: seqs[name].append(l.strip())
    return {k: "".join(v).rstrip("*") for k, v in seqs.items() if v is not None}

audit = [l.rstrip("\n").split("\t") for l in open(sys.argv[1])]
H = audit[0]; audit = [dict(zip(H, r)) for r in audit[1:]]
names_by_gene = {a["gene"]: a["name"] for a in audit}
genes = [a["gene"] for a in audit]; athal = {a["gene"]: a["best_Athal_hit"] for a in audit}

loc = {}
for l in open(V4_GFF):
    f = l.split("\t")
    if len(f) > 8 and f[2] == "gene":
        m = re.search(r"ID=gene:(\w+)", f[8])
        if m and m.group(1) in genes: loc[m.group(1)] = (f[0], int(f[3]), int(f[4]), f[6])
with open(f"{TMP}/loci.fa", "w") as o:
    for g in genes:
        c, s, e, _ = loc[g]
        o.write(subprocess.run(["samtools", "faidx", V4_GENOME, f"{c}:{s}-{e}"], capture_output=True, text=True, check=True).stdout.replace(f">{c}:{s}-{e}", f">{g}"))

def place(genome):  # best primary alignment of each locus; returns target chrom, start, end, gene strand on target
    out = subprocess.run(["minimap2", "-x", "asm5", "-t", "8", "--secondary=no", genome, f"{TMP}/loci.fa"], capture_output=True, text=True, check=True).stdout
    best = {}
    for l in out.splitlines():
        f = l.split("\t"); q, ql, st, t, ts, te, mt = f[0], int(f[1]), f[4], f[5], int(f[7]), int(f[8]), int(f[9])
        if q not in best or mt > best[q][4]:
            best[q] = (t, ts + 1, te, (loc[q][3] if st == "+" else {"+": "-", "-": "+"}[loc[q][3]]), mt, ql)
    return best

# protein models with genomic spans: (chrom, start, end, strand, protein_id)
def t2t_models():
    span = {}
    for l in open(T2T_GFF):
        f = l.split("\t")
        if len(f) > 8 and f[2] == "CDS":
            p = re.search(r"protein_id=([^;\n]+)", f[8]).group(1); s, e = int(f[3]), int(f[4])
            c0 = span.get(p, (f[0], s, e, f[6])); span[p] = (f[0], min(s, c0[1]), max(e, c0[2]), f[6])
    return [(c, s, e, st, p) for p, (c, s, e, st) in span.items()]
def refseq_models():  # NC_081805.1..NC_081823.1 = chromosomes 1..19 = CP126648.1..CP126666.1
    out = []
    for l in open(RS_PEP):
        if not l.startswith(">"): continue
        m = re.match(r">lcl\|NC_0818(\d\d)\.1_prot_(\S+)", l); lo = re.search(r"\[location=([^\]]+)\]", l)
        if not m or not lo: continue
        chrom = f"CP1266{48 + int(m.group(1)) - 5}.1"; nums = [int(x) for x in re.findall(r"\d+", lo.group(1))]
        out.append((chrom, min(nums), max(nums), "-" if "complement" in lo.group(1) else "+", l[1:].split()[0]))
    return out
def pz_models():
    out = []
    for l in open(PZ_GFF):
        f = l.split("\t")
        if len(f) > 8 and f[2] == "mRNA":
            out.append((f[0], int(f[3]), int(f[4]), f[6], re.search(r"Name=([^;\n]+)", f[8]).group(1)))
    return out

def overlapping(models, placed):
    idx = collections.defaultdict(list)
    for m in models: idx[m[0]].append(m)
    res = {}
    for g, (t, s, e, st, *_) in placed.items():
        res[g] = [m[4] for m in idx[t] if m[3] == st and m[1] <= e and m[2] >= s]
    return res

def missing(prot, at_seq):
    with open(f"{TMP}/q.faa", "w") as q: q.write(f">q\n{prot}\n")
    with open(f"{TMP}/s.faa", "w") as s: s.write(f">s\n{at_seq}\n")
    out = subprocess.run(["blastp", "-query", f"{TMP}/q.faa", "-subject", f"{TMP}/s.faa", "-evalue", "1e-6", "-outfmt",
                          "6 qstart qend qlen sstart send slen evalue bitscore"], capture_output=True, text=True).stdout.split("\n")[0]
    if not out: return None
    qs, qe, ql, ss, se, sl, *_ = out.split("\t"); qs, qe, ql, ss, se, sl = map(int, (qs, qe, ql, ss, se, sl))
    return ql, (ss - 1) - (qs - 1), (sl - se) - (ql - qe)

at = fasta(TAIR, keep=lambda n: n in set(athal.values()))
v4 = fasta(V4_PEP, keep=lambda n: n.split("_")[0] in loc and n.endswith("_P001"))
pl_t2t, pl_pz = place(T2T_GENOME), place(PZ_GENOME)
sets = {"T2T_GenBank": (overlapping(t2t_models(), pl_t2t), T2T_PEP, lambda h: h.split()[0]),
        "RefSeq": (overlapping(refseq_models(), pl_t2t), RS_PEP, lambda h: h.split()[0]),
        "Phytozome_v2.1": (overlapping(pz_models(), pl_pz), PZ_PEP, lambda h: h.split()[0])}
best = {}
for name, (ov, pep, key) in sets.items():
    want = {p for v in ov.values() for p in v}
    seqs = fasta(pep, keep=lambda n: n in want, key=key)
    for g in genes:
        cands = []
        for p in ov.get(g, []):
            if p in seqs and athal[g] in at:
                r = missing(seqs[p], at[athal[g]])
                if r and r[0] <= MAXLEN * len(at[athal[g]]): cands.append((max(r[1], 0) + max(r[2], 0), p, *r))
        best[(g, name)] = min(cands) if cands else None

# RNA-seq support: introns of the best T2T-based model (RefSeq or T2T GenBank, fewest missing residues) vs STAR junctions
# of this study's libraries mapped to the T2T genome (PN1..PN19 = CP126648.1..CP126666.1). Only the SJ.out.tab files
# present locally are used (A11, A12 = Rpv12 0 and 6 hpi); the rest are on the HPC server.
import glob
SJ = collections.Counter()
sj_files = sorted(glob.glob(R + "PN40024.T2T/STARmap/*_SJ.out.tab"))
for fn in sj_files:
    for l in open(fn):
        f = l.split("\t"); SJ[(f"CP1266{47 + int(f[0][2:])}.1", int(f[1]), int(f[2]))] += int(f[6])
introns = {}
for l in open(RS_PEP):
    if l.startswith(">lcl|NC_0818"):
        m = re.match(r">lcl\|NC_0818(\d\d)\.1_prot_(\S+)", l); lo = re.search(r"\[location=([^\]]+)\]", l)
        if m and lo:
            seg = sorted((int(a), int(b)) for a, b in re.findall(r"(\d+)\.\.(\d+)", lo.group(1)))
            introns[l[1:].split()[0]] = [(f"CP1266{48 + int(m.group(1)) - 5}.1", seg[i][1] + 1, seg[i + 1][0] - 1) for i in range(len(seg) - 1)]
cds = collections.defaultdict(list)
for l in open(T2T_GFF):
    f = l.split("\t")
    if len(f) > 8 and f[2] == "CDS": cds[re.search(r"protein_id=([^;\n]+)", f[8]).group(1)].append((f[0], int(f[3]), int(f[4])))
for pid, segs in cds.items():
    segs.sort(key=lambda x: x[1]); introns[pid] = [(segs[i][0], segs[i][2] + 1, segs[i + 1][1] - 1) for i in range(len(segs) - 1)]
def support(g):
    c = [b for b in (best[(g, "RefSeq")], best[(g, "T2T_GenBank")]) if b]
    if not c: return ["NA", "NA", "NA"]
    b = min(c); intr = introns.get(b[1], [])
    return [b[1], len(intr), sum(SJ[i] >= 3 for i in intr)]

cols = ["gene", "name", "v4_len", "Athal_hit", "v4_N_missing", "v4_C_missing", "v4_flag", "T2T_placed"]
for n in sets: cols += [f"{n}_model", f"{n}_len", f"{n}_N_missing", f"{n}_C_missing"]
cols += ["resolved_by", "T2T_best_model", "n_introns", f"introns_with_>=3_reads({len(sj_files)}_libraries)"]
print("\t".join(cols))
for a in audit:
    g = a["gene"]; t = pl_t2t.get(g); row = [g, a["name"], a["v4_len"], athal[g], a["N_missing"], a["C_missing"], a["flag"],
                                              f"{t[0]}:{t[1]}-{t[2]}{t[3]}" if t else "NA"]
    ok = []
    for n in sets:
        b = best[(g, n)]
        row += [b[1], b[2], b[3], b[4]] if b else ["NA"] * 4
        if b and a["flag"] not in ("ok", "NA") and b[3] <= MISS and b[4] <= MISS: ok.append(n)
    row.append(",".join(ok) if ok else ("-" if a["flag"] not in ("ok", "NA") else ""))
    row += support(g)
    print("\t".join(map(str, row)))

# ---- ColabFold inputs for the 10 candidate edges (Figure 4A) with the best available model per gene ----
# v4 model if not flagged; otherwise the accepted alternative (resolved_by), else the best partial T2T-based model.
# Complete precursors (signal peptides not removed). Written to gene_model_audit/corrected_models/.
import openpyxl
OUT = B + "gene_model_audit/corrected_models"
os.makedirs(OUT, exist_ok=True)
ws = list(openpyxl.load_workbook(B + "plant_immunity_Vitis/PPIs/results/Supplementary_Table_8.xlsx", read_only=True)["all_positive_hits"].iter_rows(values_only=True))
h = ws[4]; ia, ib = h.index("geneA_gene_id"), h.index("geneB_gene_id"); sec, edges = None, []
for r in ws[5:]:
    if r[3] and str(r[3]).startswith("SECTION"): sec = str(r[3]); continue
    if sec and "CANDIDATE" in sec and r[ia] and r[ib] and (r[ia], r[ib]) not in edges: edges.append((r[ia], r[ib]))
choice = {}
for a in audit:
    g = a["gene"]
    if a["flag"] in ("ok", "NA"): choice[g] = ("v4", g + "_P001", "full")
    else:
        acc = [n for n in sets if best[(g, n)] and best[(g, n)][3] <= MISS and best[(g, n)][4] <= MISS]
        pick = acc[0] if acc else min((n for n in ("RefSeq", "T2T_GenBank") if best[(g, n)]), key=lambda n: best[(g, n)], default=None)
        choice[g] = (pick, best[(g, pick)][1], "full" if acc else "partial") if pick else ("v4", g + "_P001", "truncated")
need = {g for e in edges for g in e}
seq = {}
for g in need:
    src, pid, _ = choice[g]
    path = {"v4": V4_PEP, "RefSeq": RS_PEP, "T2T_GenBank": T2T_PEP, "Phytozome_v2.1": PZ_PEP}[src]
    seq[g] = fasta(path, keep=lambda n: n == pid)[pid]
with open(f"{OUT}/network_proteins_best_models.faa", "w") as o, open(f"{OUT}/network_proteins_best_models.tsv", "w") as t:
    t.write("gene\tname\tsource\tmodel\tstatus\tlength\n")
    for g in sorted(need):
        o.write(f">{g} {choice[g][0]} {choice[g][1]} {choice[g][2]}\n{seq[g]}\n")
        t.write(f"{g}\t{names_by_gene.get(g, '') if 'names_by_gene' in globals() else ''}\t{choice[g][0]}\t{choice[g][1]}\t{choice[g][2]}\t{len(seq[g])}\n")
with open(f"{OUT}/colabfold_pairs_candidate_edges_best_models.fa", "w") as o:
    for a, b in edges:
        la, lb = (f"{g}_{names_by_gene[g]}" if names_by_gene.get(g) else g for g in (a, b))
        o.write(f">{la}:{lb}\n{seq[a]}\n:\n{seq[b]}\n")   # same layout as reference/Vitis_vinifera/GG/make_colabfold_pairs.py
print(f"# wrote {len(edges)} pairs, {len(need)} proteins to {OUT}", file=sys.stderr)
